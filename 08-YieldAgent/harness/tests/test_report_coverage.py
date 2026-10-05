import io
import asyncio
from types import SimpleNamespace
import pytest
from pptx import Presentation
from harness.types import ToolObservation, TableRef
from harness.tools import artifact_tools


def observation(result_id, tool='query_yield', scope=None, status='success', artifacts=None):
    return ToolObservation(principal_id='p', session_id='s', run_id='run', invocation_id=result_id,
        result_id=result_id, tool_name=tool, summary=result_id, scope=scope or {}, status=status,
        tables=[TableRef(table_id='data', title=result_id, columns=['product', 'value'],
            total_rows=1, preview_rows=[], complete=status == 'success', data_ref=result_id)],
        artifact_refs=artifacts or [])


def test_ordered_sections_preserve_every_source_and_analysis_status():
    from harness.report import build_report_sections
    observations = [observation('A'), observation('calc', 'run_python', {'analysis_status':'findings'}),
        observation('B'), observation('none', 'analyze_yield', {'analysis_status':'no_findings'}),
        observation('partial', status='partial')]
    tables = {o.result_id: {'data': [{'product': o.result_id, 'value': 91}]} for o in observations}
    sections = build_report_sections(observations, tables, {})
    assert list(dict.fromkeys(r for s in sections for r in s['result_ids'])) == ['A','calc','B','none','partial']
    assert [next(s['analysis_status'] for s in sections if s['result_ids'] == [r]) for r in tables] == ['not_run','findings','not_run','no_findings','partial']
    assert [s['content']['rows'][0]['product'] for s in sections if s['kind'] == 'table'] == list(tables)


def test_export_roundtrip_keeps_tables_html_tail_and_images(tmp_path, monkeypatch):
    from PIL import Image
    import ppt_builder
    monkeypatch.setattr(ppt_builder, 'OUTPUT_DIR', str(tmp_path))
    pic=io.BytesIO(); Image.new('RGB',(40,40),'red').save(pic,format='PNG')
    obs = [observation('A'), observation('calc','run_python'), observation('B'),
        observation('wads','get_wads_report',artifacts=[{'artifact_id':'html'}]),
        observation('map','render_wafer_map',artifacts=[{'artifact_id':'image'}])]
    class Store:
        async def observation(self,p,s,r): return next(o for o in obs if o.result_id==r)
        async def rows(self,p,s,r,table_id=None):
            assert table_id=='data'
            return [{'product':r,'value':91 if r=='A' else 83}]
        async def artifact(self,p,s,a):
            if a=='html': return {'mime':'text/html','title':'WADS original'}, ('<h1>Original</h1><p>'+('long body '*150)+'TAIL_EVIDENCE</p>').encode()
            return {'mime':'image/png','title':'Map'}, pic.getvalue()
    ctx=SimpleNamespace(store=Store(),run={'principal_id':'p','session_id':'s'})
    result=asyncio.run(artifact_tools.export_report(artifact_tools.ExportInput(result_ids=[o.result_id for o in obs]),ctx))
    prs=Presentation(io.BytesIO(result.artifacts[0]['data']))
    text='\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    values=[cell.text for slide in prs.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells]
    assert 'A' in values and 'B' in values and 'calc' in values and '91' in values and '83' in values
    assert 'TAIL_EVIDENCE' in text
    assert '분석 미실행' in text
    assert any(shape.shape_type==13 for slide in prs.slides for shape in slide.shapes)
    assert result.source_result_ids == [o.result_id for o in obs]
    (tmp_path/'coverage.pptx').write_bytes(result.artifacts[0]['data'])


def test_legacy_unanalyzed_summary_does_not_claim_no_findings():
    from ppt_builder import YieldReportPPTBuilder
    builder=YieldReportPPTBuilder()
    builder._add_yield_summary_slide('A',[{'week':'2026-01'}],[],'weekly')
    text='\n'.join(s.text for s in builder.prs.slides[0].shapes if s.has_text_frame)
    assert '분석 미실행' in text
    assert '이상 파라미터 없음' not in text


def test_failed_results_are_not_claimed_as_included_and_unsupported_is_explicit():
    from harness.report import build_report_sections
    failed=observation('failed',status='error')
    good=observation('good')
    sections=build_report_sections([failed,good], {'good':{'data':[{'value':1}]}}, {})
    assert {r for s in sections for r in s['result_ids']}=={'good'}
    good.artifact_refs=[{'artifact_id':'unsupported'}]
    with pytest.raises(ValueError,match='Unsupported report artifact'):
        build_report_sections([good],{'good':{'data':[{'value':1}]}}, {'unsupported':({'mime':'application/zip','title':'archive'},b'x')})


def test_report_refuses_preview_in_place_of_full_rows():
    from harness.report import build_report_sections
    obs=observation('full'); obs.tables[0].total_rows=100
    with pytest.raises(ValueError,match='Incomplete stored table'):
        build_report_sections([obs],{'full':{'data':[{'value':1}]}},{})


def test_wads_original_query_limit_marks_result_partial(monkeypatch):
    import pandas as pd
    from harness.tools.wads_tools import WadsReportInput
    import wads_tools
    frame=pd.DataFrame([dict(lotcd='A',category='PT1H',parameter='x',end_tm=f'2026-08-01 0{i}:00:00',html='<p>original</p>') for i in range(2)])
    monkeypatch.setattr(wads_tools,'_query_wads_data',lambda **kw: frame)
    async def direct(fn,*a,**kw): return fn(*a,**kw)
    monkeypatch.setattr(artifact_tools,'run_blocking',direct)
    result=asyncio.run(artifact_tools.get_wads_report(WadsReportInput(limit=1),None))
    assert result.status=='partial'
    assert not result.scope['complete']


def test_incomplete_scope_cannot_be_presented_as_no_findings():
    from harness.report import build_report_sections
    obs=observation('partial-scope',scope={'complete':False,'analysis_status':'no_findings'})
    sections=build_report_sections([obs],{obs.result_id:{'data':[{'value':1}]}},{})
    assert all(s['analysis_status']=='partial' for s in sections)


def test_korean_text_pagination_accounts_for_wide_glyphs(tmp_path,monkeypatch):
    import unicodedata
    import ppt_builder
    monkeypatch.setattr(ppt_builder,'OUTPUT_DIR',str(tmp_path))
    section={'kind':'text','title':'Korean','result_ids':['r'],'table_ids':[], 'analysis_status':'not_run','content':{'text':'가나다라마바사'*50}}
    payload,_=ppt_builder.YieldReportPPTBuilder().build({'report_sections':[section]})
    prs=Presentation(io.BytesIO(payload))
    body=[shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame and '가나다' in shape.text]
    assert body
    assert all(sum(2 if unicodedata.east_asian_width(c) in ('W','F') else 1 for c in line)<=85 for text in body for line in text.splitlines())
